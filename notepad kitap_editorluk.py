#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os, sys, re, time
from pathlib import Path

def kur(paket, import_adi=None):
    try:
        __import__(import_adi or paket)
    except ImportError:
        print(f"  {paket} kuruluyor...")
        os.system(f"{sys.executable} -m pip install {paket} -q")

print("Kontrol ediliyor...")
kur("requests")
kur("python-docx", "docx")
kur("pymupdf", "fitz")
print("Hazir.\n")

import requests

def oku_pdf(yol):
    import fitz
    doc = fitz.open(yol)
    metin = "\n\n".join(sayfa.get_text() for sayfa in doc)
    doc.close()
    return metin

def oku_txt(yol):
    for enc in ("utf-8", "utf-8-sig", "cp1254", "latin-1"):
        try:
            return Path(yol).read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("Dosya okunamadi.")

def oku_docx(yol):
    import docx
    doc = docx.Document(yol)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

def temizle(metin):
    # NULL ve kontrol karakterlerini temizle
    metin = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', metin)
    metin = re.sub(r'\ufffd', '', metin)
    return metin

def github_cek(url):
    raw = url
    if "github.com" in url and "raw.githubusercontent.com" not in url:
        raw = url.replace("github.com", "raw.githubusercontent.com").replace("/blob/", "/")
    r = requests.get(raw, timeout=30)
    r.raise_for_status()
    # PDF ise geçici dosyaya indir
    if url.endswith(".pdf"):
        tmp = Path("_tmp_kitap.pdf")
        tmp.write_bytes(r.content)
        metin = oku_pdf(str(tmp))
        tmp.unlink()
        return metin
    return r.text

def kaynak_yukle(kaynak):
    if kaynak.startswith("http"):
        print("GitHub'dan cekiliyor...")
        metin = github_cek(kaynak)
        ad = "kitap"
        return temizle(metin), ad
    p = Path(kaynak)
    if not p.exists():
        raise FileNotFoundError(f"Bulunamadi: {kaynak}")
    ext = p.suffix.lower()
    okuyucular = {".txt": oku_txt, ".docx": oku_docx, ".pdf": oku_pdf}
    if ext not in okuyucular:
        raise ValueError(f"Desteklenmeyen format: {ext}")
    return temizle(okuyucular[ext](kaynak)), p.stem

YAZIM_DUZELTME = [
    (r"\b(\w{3,})\s+\1\b", r"\1"),
    (r"  +", " "),
    (r"\s+([,\.!?:;])", r"\1"),
    (r"([,\.!?:;])([^\s\d\n])", r"\1 \2"),
    (r"\(\s+", "("),
    (r"\s+\)", ")"),
]

def yazim_duzelt(metin):
    for pattern, repl in YAZIM_DUZELTME:
        metin = re.sub(pattern, repl, metin, flags=re.IGNORECASE)
    return metin

def noktalama_duzelt(metin):
    def buyuk_harf(m):
        return m.group(1) + m.group(2).upper()
    return re.sub(r"([.!?]\s+)([a-z])", buyuk_harf, metin)

def paragraf_duzelt(metin):
    satirlar = metin.split("\n")
    paragraflar, tampon = [], []
    for satir in satirlar:
        satir = satir.strip()
        if not satir:
            if tampon:
                paragraflar.append(" ".join(tampon))
                tampon = []
        else:
            tampon.append(satir)
    if tampon:
        paragraflar.append(" ".join(tampon))
    return "\n\n".join(p for p in paragraflar if p)

def tekrar_temizle(metin):
    return re.sub(r"\b(\w{3,})\s+\1\b", r"\1", metin, flags=re.IGNORECASE)

def editorluk_yap(metin):
    adimlar = [
        ("Yazim duzeltiliyor",     yazim_duzelt),
        ("Noktalama duzeltiliyor", noktalama_duzelt),
        ("Tekrarlar temizleniyor", tekrar_temizle),
        ("Paragraf yapisi",        paragraf_duzelt),
    ]
    for aciklama, fonk in adimlar:
        print(f"  {aciklama}...")
        time.sleep(0.2)
        metin = fonk(metin)
    return metin

def kaydet_txt(metin, kaynak_adi):
    yol = Path(f"{kaynak_adi}_duzenlemis.txt")
    yol.write_text(metin, encoding="utf-8")
    return yol

def kaydet_docx(metin, kaynak_adi):
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    yol = Path(f"{kaynak_adi}_duzenlemis.docx")
    doc = docx.Document()
    for bolum in doc.sections:
        bolum.left_margin = Cm(3)
        bolum.right_margin = Cm(2.5)
    for para in metin.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        # XML uyumsuz karakterleri temizle
        para = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', para)
        para = ''.join(c for c in para if c.isprintable() or c in '\n\t')
        if not para:
            continue
        try:
            p = doc.add_paragraph(para)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if p.runs:
                p.runs[0].font.size = Pt(12)
            p.paragraph_format.first_line_indent = Cm(1)
        except Exception:
            continue
    doc.save(str(yol))
    return yol

def main():
    print("=" * 50)
    print("   KITAP EDITORLUK SISTEMI  (API'siz)")
    print("=" * 50)

    print("\nKaynak secin:")
    print("  1  Bilgisayardan dosya (.txt / .docx / .pdf)")
    print("  2  GitHub URL")
    secim = input("Secim (1/2): ").strip()

    if secim == "1":
        kaynak = input("Dosya yolu: ").strip().strip('"')
    else:
        kaynak = input("GitHub URL: ").strip()

    try:
        metin, kaynak_adi = kaynak_yukle(kaynak)
    except Exception as e:
        print(f"Hata: {e}")
        sys.exit(1)

    print(f"\nYuklendi: {len(metin):,} karakter")

    print("\nCikti formati:")
    print("  1  .txt")
    print("  2  .docx")
    fmt = input("Secim (1/2): ").strip()

    print("\nEditorluk basliyor...\n")
    duzenlemis = editorluk_yap(metin)

    if fmt == "2":
        cikti = kaydet_docx(duzenlemis, kaynak_adi)
    else:
        cikti = kaydet_txt(duzenlemis, kaynak_adi)

    print(f"\nTamamlandi! Kaydedildi: {cikti.resolve()}")
    print("\nOnizleme (ilk 300 karakter):")
    print("-" * 40)
    print(duzenlemis[:300])

if __name__ == "__main__":
    main()